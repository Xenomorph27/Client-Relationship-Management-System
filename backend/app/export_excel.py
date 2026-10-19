"""Export every table in the database to one Excel workbook - a sheet per table.

Run from this directory (backend/app), because the project uses flat imports:

    ../../.venv/Scripts/python.exe export_excel.py
    ../../.venv/Scripts/python.exe export_excel.py --out somewhere/else.xlsx

Read-only: it runs SELECT statements and never modifies the database.

Every cell is a plain value rather than a formula, because this is a snapshot of
stored data - there is nothing here to recalculate.
"""
import sys
from datetime import datetime, date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from database.database import SessionLocal

DEFAULT_OUT = "../../crm_database_export.xlsx"

# Sheet title -> table name. Parents first so the workbook reads top-down.
TABLES = [
    ("Users", "users"),
    ("Organisations", "organisations"),
    ("Leads", "leads"),
    ("Customers", "customers"),
    ("Deals", "deals"),
    ("Emails", "emails"),
    ("Campaigns", "campaigns"),
]

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
BODY_FONT = Font(name="Arial")
TITLE_FONT = Font(name="Arial", bold=True, size=14)

MAX_WIDTH = 55
DATE_FMT = "%Y-%m-%d %H:%M"


def cell_value(value):
    """Convert a database value into something Excel displays sensibly."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.strftime(DATE_FMT)
    if hasattr(value, "value"):  # a Python Enum, e.g. Priority_Level
        return value.value
    return value


def write_sheet(ws, columns, rows):
    ws.append(list(columns))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append([cell_value(v) for v in row])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT

    # Width from the longest value in each column, capped so long email
    # bodies do not push a column off the screen.
    for i, column in enumerate(columns, start=1):
        longest = len(str(column))
        for row in rows:
            longest = max(longest, len(str(cell_value(row[i - 1]))))
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 3, MAX_WIDTH)

    ws.freeze_panes = "A2"  # keep headers visible while scrolling
    if rows:
        ws.auto_filter.ref = ws.dimensions  # dropdown filters on every column


def write_overview(ws, counts, exported_at):
    ws["A1"] = "CRM Database Export"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Exported {exported_at.strftime('%Y-%m-%d %H:%M')} from database 'crm_db'"
    ws["A2"].font = Font(name="Arial", italic=True, size=10)

    ws["A4"] = "Sheet"
    ws["B4"] = "Table"
    ws["C4"] = "Rows"
    for cell in (ws["A4"], ws["B4"], ws["C4"]):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    row = 5
    for title, table in TABLES:
        ws.cell(row=row, column=1, value=title).font = BODY_FONT
        ws.cell(row=row, column=2, value=table).font = BODY_FONT
        ws.cell(row=row, column=3, value=counts[table]).font = BODY_FONT
        row += 1

    ws.cell(row=row, column=1, value="Total").font = Font(name="Arial", bold=True)
    ws.cell(row=row, column=3, value=sum(counts.values())).font = Font(name="Arial", bold=True)

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 10


def main():
    out = DEFAULT_OUT
    if "--out" in sys.argv:
        out = sys.argv[sys.argv.index("--out") + 1]

    db = SessionLocal()
    wb = Workbook()
    counts = {}
    try:
        overview = wb.active
        overview.title = "Overview"

        for title, table in TABLES:
            result = db.execute(text(f"SELECT * FROM {table} ORDER BY id"))
            columns = list(result.keys())
            rows = result.fetchall()
            counts[table] = len(rows)
            write_sheet(wb.create_sheet(title), columns, rows)
            print(f"{title}: {len(rows)} rows, {len(columns)} columns")

        write_overview(overview, counts, datetime.now())
        wb.save(out)
        print(f"\nsaved {out} ({sum(counts.values())} rows across {len(TABLES)} sheets)")
    finally:
        db.close()


if __name__ == "__main__":
    main()
